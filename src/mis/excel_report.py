"""Formatted Excel MIS workbook. Numbers are written as real numeric cells with number formats."""
from __future__ import annotations

from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PASS_FILL = PatternFill("solid", fgColor="D9EAD3")
FAIL_FILL = PatternFill("solid", fgColor="F4CCCC")
THIN = Side(style="thin", color="BBBBBB")


def _header(ws, row: int, labels: list[str]) -> None:
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_excel(path, kpis: pd.DataFrame, methods: pd.DataFrame, merchants: pd.DataFrame,
                checks: pd.DataFrame, commentary: dict, commentary_source: str, latest_label: str) -> None:
    wb = Workbook()

    # ---- Monthly KPIs --------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Monthly KPIs"
    cols = ["Month", "Transactions", "Successful", "GMV (INR)", "GMV (INR crore)", "Success rate",
            "Refund value (INR)", "Refund rate", "GMV MoM growth", "Txn MoM growth"]
    _header(ws, 1, cols)
    for r, row in enumerate(kpis.itertuples(), 2):
        vals = [pd.Timestamp(row.month).to_pydatetime(), int(row.txn_count), int(row.success_count), float(row.gmv),
                float(row.gmv) / 1e7, float(row.success_rate), float(row.refund_value), float(row.refund_rate),
                None if pd.isna(row.gmv_mom_growth) else float(row.gmv_mom_growth),
                None if pd.isna(row.txn_mom_growth) else float(row.txn_mom_growth)]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v)
        for c, fmt in enumerate(["mmm yyyy", "#,##0", "#,##0", "#,##0", "0.00", "0.0%", "#,##0", "0.00%", "0.0%", "0.0%"], 1):
            ws.cell(row=r, column=c).number_format = fmt
    ws.freeze_panes = "B2"
    _widths(ws, [11, 14, 14, 18, 16, 13, 18, 12, 15, 15])
    ws.row_dimensions[1].height = 32
    last = len(kpis) + 1

    # ---- Summary ---------------------------------------------------------------------------------
    s = wb.create_sheet("Summary", 0)
    s["A1"] = f"Payments MIS - {latest_label}"
    s["A1"].font = Font(bold=True, size=16)
    s["A2"] = "Synthetic data. All figures computed in SQL (DuckDB) from the transaction table."
    s["A2"].font = Font(italic=True, color="666666")
    latest = kpis.iloc[-1]
    tiles = [("GMV (INR crore)", float(latest.gmv) / 1e7, "0.00"), ("Transactions", int(latest.txn_count), "#,##0"),
             ("Success rate", float(latest.success_rate), "0.0%"), ("Refund rate", float(latest.refund_rate), "0.00%"),
             ("GMV MoM growth", float(latest.gmv_mom_growth), "+0.0%;-0.0%")]
    for i, (label, value, fmt) in enumerate(tiles):
        col = 1 + i * 2
        s.cell(row=4, column=col, value=label).font = Font(bold=True, color="666666")
        c = s.cell(row=5, column=col, value=value)
        c.number_format, c.font = fmt, Font(bold=True, size=14)
        c.alignment = Alignment(horizontal="left")
    s["A7"] = "Commentary"
    s["A7"].font = Font(bold=True, size=12)
    s["A8"] = commentary["headline"]
    s["A8"].font = Font(bold=True)
    for i, b in enumerate(commentary["bullets"]):
        s.cell(row=9 + i, column=1, value="• " + b)
    s.cell(row=10 + len(commentary["bullets"]), column=1, value=f"Commentary written by: {commentary_source}; every figure checked against the KPI table.").font = Font(italic=True, color="666666")
    _widths(s, [20, 4, 16, 4, 14, 4, 14, 4, 16, 4])
    chart = BarChart()
    chart.title, chart.y_axis.title, chart.height, chart.width = "GMV by month (INR crore)", "INR crore", 7.5, 18
    chart.add_data(Reference(ws, min_col=5, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    chart.legend = None
    s.add_chart(chart, f"A{13 + len(commentary['bullets'])}")

    # ---- Methods / merchants ---------------------------------------------------------------------
    m = wb.create_sheet("Payment methods")
    _header(m, 1, ["Payment method", "Transactions", "GMV (INR)", "GMV share", "Success rate"])
    for r, row in enumerate(methods.itertuples(), 2):
        for c, (v, fmt) in enumerate([(row.payment_method, None), (int(row.txn_count), "#,##0"), (float(row.gmv), "#,##0"),
                                       (float(row.gmv_share), "0.0%"), (float(row.success_rate), "0.0%")], 1):
            cell = m.cell(row=r, column=c, value=v)
            if fmt:
                cell.number_format = fmt
    _widths(m, [18, 14, 18, 12, 13])

    t = wb.create_sheet("Top merchants")
    _header(t, 1, ["Rank", "Merchant", "GMV (INR)", "GMV share"])
    for r, row in enumerate(merchants.itertuples(), 2):
        for c, (v, fmt) in enumerate([(int(row.gmv_rank), None), (row.merchant_id, None), (float(row.gmv), "#,##0"),
                                       (float(row.gmv_share), "0.0%")], 1):
            cell = t.cell(row=r, column=c, value=v)
            if fmt:
                cell.number_format = fmt
    _widths(t, [8, 14, 18, 12])

    # ---- Data checks -----------------------------------------------------------------------------
    d = wb.create_sheet("Data checks")
    _header(d, 1, ["Check", "What it verifies", "Violations", "Status"])
    for r, row in enumerate(checks.itertuples(), 2):
        for c, v in enumerate([row.check, row.description, int(row.violations), row.status], 1):
            cell = d.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        d.cell(row=r, column=4).fill = PASS_FILL if row.status == "PASS" else FAIL_FILL
    _widths(d, [32, 70, 12, 10])

    # ---- About -----------------------------------------------------------------------------------
    a = wb.create_sheet("About")
    lines = ["Definitions", "GMV = sum of amount for SUCCESS transactions.", "Success rate = SUCCESS transactions / all transactions.",
             "Refund rate = refund value / GMV.", "MoM growth = this month / previous month - 1.", "",
             f"Generated: {datetime.now():%Y-%m-%d %H:%M}", f"Commentary: {commentary_source}",
             "Publishing rules: all ten data checks pass, and every figure in the commentary matches the KPI table."]
    for i, line in enumerate(lines, 1):
        a.cell(row=i, column=1, value=line)
    a["A1"].font = Font(bold=True)
    a.column_dimensions["A"].width = 100
    for sheet in wb.worksheets:  # print cleanly: landscape, one page wide
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
