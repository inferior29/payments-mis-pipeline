import pandas as pd
import pytest
from openpyxl import load_workbook

from mis.commentary import CommentaryMismatch, TemplateCommentary
from mis.generate_data import generate
from mis.run import build_report
from conftest import make_con


class FixedCommentary:
    name = "fixed"

    def __init__(self, commentary):
        self._c = commentary

    def write(self, ctx):
        return self._c


def test_generator_is_deterministic_and_well_formed():
    a, b = generate(5000, seed=3), generate(5000, seed=3)
    assert a.equals(b) and a["txn_id"].is_unique and len(a) == 5000
    assert set(a["status"]) <= {"SUCCESS", "FAILED", "PENDING"}
    assert (a.loc[a.refund_amount > 0, "status"] == "SUCCESS").all()


def test_report_is_written_with_real_numeric_cells(big_df, tmp_path):
    con = make_con(big_df)
    path = build_report(con, TemplateCommentary(), tmp_path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Monthly KPIs", "Payment methods", "Top merchants", "Data checks", "About"]
    k = wb["Monthly KPIs"]
    assert k.max_row == 13                                               # header + 12 months
    gmv_cell, rate_cell = k["D13"], k["F13"]
    assert isinstance(gmv_cell.value, float) and gmv_cell.number_format == "#,##0"
    assert rate_cell.number_format == "0.0%"
    sep = big_df[(big_df.status == "SUCCESS") & (big_df.txn_ts >= "2026-09-01")]["amount"].sum()
    assert gmv_cell.value == pytest.approx(sep, rel=1e-9)               # workbook equals an independent pandas total
    assert [r[3] for r in wb["Data checks"].iter_rows(min_row=2, values_only=True)] == ["PASS"] * 10


def test_report_is_blocked_when_a_data_check_fails(big_df, tmp_path):
    bad = big_df.copy()
    bad.loc[0, "amount"] = -1.0
    with pytest.raises(RuntimeError, match="Data checks failed"):
        build_report(make_con(bad), TemplateCommentary(), tmp_path)
    assert list(tmp_path.glob("*.xlsx")) == []


def test_report_is_blocked_when_commentary_has_a_wrong_figure(big_df, tmp_path):
    con = make_con(big_df)
    good = TemplateCommentary().write(__import__("mis.kpis", fromlist=["build_context"]).build_context(con))
    good["headline"] = good["headline"].replace("₹", "₹9")             # corrupt one figure
    with pytest.raises(CommentaryMismatch):
        build_report(con, FixedCommentary(good), tmp_path)
    assert list(tmp_path.glob("*.xlsx")) == []                           # nothing was published
